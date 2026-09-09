import json
import os
import sys
import types
from pathlib import Path
from unittest import mock

import pytest

from test_support import authenticated_mcp_params, load_local_module


pytestmark = pytest.mark.skipif(
    sys.platform != "linux" or not hasattr(os, "O_NOFOLLOW"),
    reason="the document app's descriptor-safe contract targets Claw OS Linux",
)


@pytest.fixture
def doc_app():
    return load_local_module(Path(__file__).with_name("main.py"), "doc_app_main")


@pytest.fixture
def sdk_app(monkeypatch, doc_app):
    monkeypatch.setenv("COS_APP_MANIFEST", str(Path(__file__).with_name("app.json")))
    monkeypatch.setenv("COS_APP_ID", "doc")
    with mock.patch.dict(sys.modules, {"main": doc_app}), mock.patch(
        "claw_os_sdk.mcp.App.serve", autospec=True,
    ) as serve:
        load_local_module(Path(__file__).with_name("server.py"), "doc_app_server")
    serve.assert_called_once()
    return serve.call_args.args[0]


def _call(app, command, arguments):
    return app._handle_request(
        "tools/call",
        authenticated_mcp_params({"name": f"doc.{command}", "arguments": arguments}),
        True,
    )


def _ai_response(text="Presented document"):
    return types.SimpleNamespace(
        text=text, model="fixture-model", provider="fixture-provider",
        usage=types.SimpleNamespace(input_tokens=10, output_tokens=5, units=15),
        budget=types.SimpleNamespace(period="fixture", units_used=15, units_cap=200000),
        review=types.SimpleNamespace(safety="strict", prompt_redacted=False),
    )


def test_original_manifest_authority_and_sdk_catalog_are_preserved(sdk_app):
    manifest = json.loads(Path(__file__).with_name("app.json").read_text())
    assert manifest["id"] == "doc"
    assert manifest["ai"] == {
        "budget": {"monthly_units": 200000}, "safety": "strict",
        "origins": ["external-content"],
    }
    operations = manifest["operations"]
    assert list(operations) == ["read", "info", "convert", "summarize", "explain", "rewrite"]
    assert operations["read"]["needs"][0]["verb"] == "fs.read"
    assert operations["info"]["needs"][0]["verb"] == "fs.meta"
    assert operations["convert"]["needs"][1]["scope"] == {"kind": "wild"}
    assert operations["summarize"]["needs"][2]["scope"]["scope"] == {
        "kind": "self-ref", "value": "doc",
    }
    assert [tool["name"] for tool in sdk_app._handle_request("tools/list", {}, True)["tools"]] == [
        f"doc.{name}" for name in operations
    ]
    for tool in manifest["mcp"]["tools"]:
        operation = operations[tool["name"].removeprefix("doc.")]
        assert tool["needs"] == operation["needs"]
        assert tool["args"] == [
            {key: value for key, value in arg.items() if key != "binding"}
            for arg in operation["args"]
        ]


def test_sdk_dispatches_read_info_and_convert_with_exact_scopes(sdk_app, doc_app, tmp_path):
    source = tmp_path / "records.json"
    source.write_text('[{"value":"one"}]')
    output = tmp_path / "records.csv"
    with mock.patch.object(doc_app.policy, "require") as require:
        read = _call(sdk_app, "read", {"path": str(source)})
        assert json.loads(read["structuredContent"]["content"]) == [{"value": "one"}]
        require.assert_called_once_with("fs.read", path=str(source))
        require.reset_mock()
        info = _call(sdk_app, "info", {"path": str(source)})
        assert info["structuredContent"] == {
            "path": str(source), "format": "json", "size": source.stat().st_size, "readable": True,
        }
        require.assert_called_once_with("fs.meta", path=str(source))
        require.reset_mock()
        convert = _call(sdk_app, "convert", {"path": str(source), "to": "csv"})
        assert convert["structuredContent"] == {
            "input": str(source), "output": str(output), "format": "csv",
        }
        assert output.read_text() == "value\none\n"
        assert require.call_args_list == [
            mock.call("fs.read", path=str(source)), mock.call("fs.write", path=str(output)),
        ]


@pytest.mark.parametrize(("command", "arguments"), [
    ("read", {}), ("info", {"path": 5}), ("convert", {"path": "/unused"}),
    ("rewrite", {"text": ["safe"], "session_id": "forged"}),
    ("summarize", {"text": "not-a-repeatable-array"}),
])
def test_sdk_rejects_invalid_arguments_before_authority(sdk_app, doc_app, command, arguments):
    with mock.patch.object(doc_app.policy, "require") as require, mock.patch.object(
        doc_app.ai, "chat",
    ) as chat, mock.patch.object(doc_app.memory, "remember") as remember:
        assert _call(sdk_app, command, arguments)["isError"]
        require.assert_not_called()
        chat.assert_not_called()
        remember.assert_not_called()


@pytest.mark.parametrize(("command", "max_units"), [
    ("summarize", 6000), ("explain", 4000), ("rewrite", 8000),
])
def test_sdk_ai_uses_untrusted_document_and_doc_memory_identity(
    sdk_app, doc_app, tmp_path, command, max_units,
):
    source = tmp_path / "document.txt"
    source.write_text("untrusted document text")
    arguments = {"file": str(source)}
    if command == "rewrite":
        arguments["instruction"] = "--keep-meaning"
    with mock.patch.object(doc_app.policy, "require") as require, mock.patch.object(
        doc_app.ai, "chat", return_value=_ai_response(),
    ) as chat, mock.patch.object(doc_app.memory, "remember") as remember:
        result = _call(sdk_app, command, arguments)["structuredContent"]
        assert require.call_args_list == [
            mock.call("fs.read", path=str(source)), mock.call("ai.chat.untrusted", wild=True),
        ]
        assert chat.call_args.kwargs["prompt"] == "untrusted document text"
        assert chat.call_args.kwargs["origin"] == "external-content"
        assert chat.call_args.kwargs["max_units"] == max_units
        assert result["source"] == str(source)
        assert result["model"] == "fixture-model"
        assert result["provider"] == "fixture-provider"
        assert result["usage"] == {"input_tokens": 10, "output_tokens": 5, "units": 15}
        assert result["budget"] == {"period": "fixture", "units_used": 15, "units_cap": 200000}
        assert source.read_text() == "untrusted document text"
        if command == "summarize":
            assert result["summary"] == "Presented document"
            remember.assert_called_once_with(
                source="doc", text=f"Summarised document {source}: Presented document",
                kind="note", entity_id=str(source), tags=["doc", "summary"],
                link=f"cos doc summarize --file {source}",
            )
        else:
            assert result["text"] == "Presented document"
            remember.assert_not_called()
        if command == "rewrite":
            assert chat.call_args.kwargs["system"].endswith("Instruction: --keep-meaning")


def test_sdk_denial_stops_ai_and_memory_without_exposing_content(sdk_app, doc_app, tmp_path):
    source = tmp_path / "document.txt"
    source.write_text("private document text")
    with mock.patch.object(
        doc_app.policy, "require",
        side_effect=doc_app.policy.PermissionDenied({"summary": "denied"}),
    ), mock.patch.object(doc_app.ai, "chat") as chat, mock.patch.object(
        doc_app.memory, "remember",
    ) as remember:
        result = _call(sdk_app, "summarize", {"file": str(source)})
        assert result["isError"]
        assert result["structuredContent"]["denial"] == {"summary": "denied"}
        assert "private document text" not in json.dumps(result)
        chat.assert_not_called()
        remember.assert_not_called()


def test_ai_bounds_input_and_does_not_remember_inline_presentation(sdk_app, doc_app):
    with mock.patch.object(doc_app.policy, "require") as require, mock.patch.object(
        doc_app.ai, "chat", return_value=_ai_response(),
    ) as chat, mock.patch.object(doc_app.memory, "remember") as remember:
        result = _call(sdk_app, "summarize", {"text": ["x" * 100001]})["structuredContent"]
        assert len(chat.call_args.kwargs["prompt"]) == 100000
        assert result["source_chars"] == 100000
        assert result["source"] is None
        require.assert_called_once_with("ai.chat.untrusted", wild=True)
        remember.assert_not_called()


def _swap_during_authorization(
    monkeypatch,
    doc_app,
    requested,
    replacement,
    *,
    verb="fs.read",
):
    authorized_path = os.path.realpath(requested)

    def require(actual_verb, *, path=None, **_scope):
        assert actual_verb == verb
        assert path == authorized_path
        requested.unlink()
        requested.symlink_to(replacement)

    monkeypatch.setattr(doc_app.policy, "require", require)


def test_read_keeps_opened_file_when_path_becomes_symlink(
    tmp_path,
    monkeypatch,
    doc_app,
):
    requested = tmp_path / "note.txt"
    replacement = tmp_path / "secret.txt"
    requested.write_text("authorized content", encoding="utf-8")
    replacement.write_text("secret content", encoding="utf-8")
    _swap_during_authorization(
        monkeypatch,
        doc_app,
        requested,
        replacement,
    )

    result = doc_app.cmd_read([str(requested)])

    assert result["content"] == "authorized content"
    assert requested.read_text(encoding="utf-8") == "secret content"


def test_read_rejects_final_symlink_before_policy(
    tmp_path,
    monkeypatch,
    doc_app,
):
    secret = tmp_path / "secret.txt"
    link = tmp_path / "link.txt"
    secret.write_text("secret content", encoding="utf-8")
    link.symlink_to(secret)

    def unexpected_authorization(*_args, **_kwargs):
        pytest.fail("a final symlink must be rejected before authorization")

    monkeypatch.setattr(doc_app.policy, "require", unexpected_authorization)

    result = doc_app.cmd_read([str(link)])

    assert "symlink" in result["error"]
    assert "secret content" not in result["error"]


def test_info_uses_fstat_from_authorized_descriptor(
    tmp_path,
    monkeypatch,
    doc_app,
):
    requested = tmp_path / "report.unknown"
    replacement = tmp_path / "secret.unknown"
    requested.write_text("short", encoding="utf-8")
    replacement.write_text("secret content that is longer", encoding="utf-8")
    _swap_during_authorization(
        monkeypatch,
        doc_app,
        requested,
        replacement,
        verb="fs.meta",
    )

    result = doc_app.cmd_info([str(requested)])

    assert result["size"] == len("short")
    assert result["readable"] is True


@pytest.mark.parametrize(
    ("suffix", "authorized", "secret"),
    [
        (".txt", "authorized text", "secret text"),
        (".md", "authorized markdown", "secret markdown"),
        (".json", '{"value": "authorized json"}', '{"value": "secret json"}'),
        (".csv", "value\nauthorized csv\n", "value\nsecret csv\n"),
        (".yaml", "value: authorized yaml\n", "value: secret yaml\n"),
        (".unknown", "authorized fallback", "secret fallback"),
    ],
)
def test_text_readers_share_authorized_descriptor_contract(
    tmp_path,
    monkeypatch,
    doc_app,
    suffix,
    authorized,
    secret,
):
    requested = tmp_path / f"document{suffix}"
    replacement = tmp_path / f"secret{suffix}"
    requested.write_text(authorized, encoding="utf-8")
    replacement.write_text(secret, encoding="utf-8")
    _swap_during_authorization(
        monkeypatch,
        doc_app,
        requested,
        replacement,
    )

    result = doc_app.cmd_read([str(requested)])

    assert "authorized" in result["content"]
    assert "secret" not in result["content"]


def test_packaged_readers_share_authorized_descriptor_contract(
    tmp_path,
    monkeypatch,
    doc_app,
):
    authorized = b"authorized package bytes"
    secret = b"secret package bytes"

    class FakePdf:
        def __iter__(self):
            return iter([types.SimpleNamespace(get_text=lambda: "authorized pdf")])

        def close(self):
            return None

    def open_pdf(*, stream, filetype):
        assert stream == authorized
        assert filetype == "pdf"
        return FakePdf()

    monkeypatch.setitem(sys.modules, "fitz", types.SimpleNamespace(open=open_pdf))

    def open_docx(source):
        assert source.read() == authorized
        return types.SimpleNamespace(
            paragraphs=[types.SimpleNamespace(text="authorized docx")]
        )

    monkeypatch.setitem(
        sys.modules,
        "docx",
        types.SimpleNamespace(Document=open_docx),
    )

    class FakeWorkbook:
        sheetnames = ["Sheet1"]

        def __getitem__(self, _name):
            return types.SimpleNamespace(
                iter_rows=lambda **_kwargs: iter([("authorized xlsx",)])
            )

        def close(self):
            return None

    def open_xlsx(source, *, read_only, data_only):
        assert source.read() == authorized
        assert read_only is True
        assert data_only is True
        return FakeWorkbook()

    monkeypatch.setitem(
        sys.modules,
        "openpyxl",
        types.SimpleNamespace(load_workbook=open_xlsx),
    )

    def open_pptx(source):
        assert source.read() == authorized
        paragraph = types.SimpleNamespace(text="authorized pptx")
        shape = types.SimpleNamespace(
            has_text_frame=True,
            text_frame=types.SimpleNamespace(paragraphs=[paragraph]),
        )
        slide = types.SimpleNamespace(
            shapes=[shape],
            has_notes_slide=False,
        )
        return types.SimpleNamespace(slides=[slide])

    monkeypatch.setitem(
        sys.modules,
        "pptx",
        types.SimpleNamespace(Presentation=open_pptx),
    )

    for suffix in (".pdf", ".docx", ".xlsx", ".pptx"):
        requested = tmp_path / f"document{suffix}"
        replacement = tmp_path / f"secret{suffix}"
        requested.write_bytes(authorized)
        replacement.write_bytes(secret)
        _swap_during_authorization(
            monkeypatch,
            doc_app,
            requested,
            replacement,
        )

        result = doc_app.cmd_read([str(requested)])

        assert "authorized" in result["content"]
        assert "secret" not in result["content"]


def test_canonical_inline_flag_and_delimited_text_reach_handler(monkeypatch, doc_app):
    monkeypatch.setattr(
        doc_app.sys,
        "stdin",
        types.SimpleNamespace(
            isatty=lambda: False,
            read=lambda: "piped input must not replace positional text",
        ),
    )
    monkeypatch.setattr(doc_app, "_ai_call", lambda **kwargs: kwargs)

    result = doc_app.run(
        "rewrite",
        ["--instruction=--urgent", "--", "--file"],
    )

    assert result["text"] == "--file"
    assert result["source"] is None


def test_real_packaged_readers_survive_symlink_swap(
    tmp_path,
    monkeypatch,
    doc_app,
):
    fitz = pytest.importorskip("fitz")
    docx = pytest.importorskip("docx")
    openpyxl = pytest.importorskip("openpyxl")
    pptx = pytest.importorskip("pptx")
    from pptx.util import Inches

    def write_pdf(path, text):
        document = fitz.open()
        page = document.new_page()
        page.insert_text((72, 72), text)
        document.save(path)
        document.close()

    def write_docx(path, text):
        document = docx.Document()
        document.add_paragraph(text)
        document.save(path)

    def write_xlsx(path, text):
        workbook = openpyxl.Workbook()
        workbook.active["A1"] = text
        workbook.save(path)
        workbook.close()

    def write_pptx(path, text):
        presentation = pptx.Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        box = slide.shapes.add_textbox(
            Inches(1),
            Inches(1),
            Inches(5),
            Inches(1),
        )
        box.text = text
        presentation.save(path)

    writers = {
        ".pdf": write_pdf,
        ".docx": write_docx,
        ".xlsx": write_xlsx,
        ".pptx": write_pptx,
    }
    for suffix, writer in writers.items():
        requested = tmp_path / f"real{suffix}"
        replacement = tmp_path / f"secret-real{suffix}"
        writer(requested, f"authorized {suffix}")
        writer(replacement, f"secret {suffix}")
        _swap_during_authorization(
            monkeypatch,
            doc_app,
            requested,
            replacement,
        )

        result = doc_app.cmd_read([str(requested)])

        assert "authorized" in result["content"]
        assert "secret" not in result["content"]


def test_convert_reads_opened_source_after_symlink_swap(
    tmp_path,
    monkeypatch,
    doc_app,
):
    requested = tmp_path / "records.json"
    replacement = tmp_path / "secret.json"
    requested.write_text('[{"value": "authorized"}]', encoding="utf-8")
    replacement.write_text('[{"value": "secret"}]', encoding="utf-8")
    authorized_path = os.path.realpath(requested)
    swapped = False

    def require(verb, *, path=None, **_scope):
        nonlocal swapped
        if verb == "fs.read":
            assert path == authorized_path
            requested.unlink()
            requested.symlink_to(replacement)
            swapped = True
        else:
            assert verb == "fs.write"

    monkeypatch.setattr(doc_app.policy, "require", require)

    result = doc_app.cmd_convert([str(requested), "--to", "csv"])

    assert swapped is True
    assert result["format"] == "csv"
    output = Path(result["output"]).read_text(encoding="utf-8")
    assert "authorized" in output
    assert "secret" not in output


def test_convert_rejects_output_symlink_after_authorization(
    tmp_path,
    monkeypatch,
    doc_app,
):
    source = tmp_path / "records.json"
    outside = tmp_path / "outside.csv"
    output = tmp_path / "records.csv"
    source.write_text('[{"value": "authorized"}]', encoding="utf-8")
    outside.write_text("do not overwrite", encoding="utf-8")
    output.symlink_to(outside)
    monkeypatch.setattr(doc_app.policy, "require", lambda *_args, **_kwargs: None)

    result = doc_app.cmd_convert([str(source), "--to", "csv"])

    assert "output symlink" in result["error"]
    assert outside.read_text(encoding="utf-8") == "do not overwrite"
